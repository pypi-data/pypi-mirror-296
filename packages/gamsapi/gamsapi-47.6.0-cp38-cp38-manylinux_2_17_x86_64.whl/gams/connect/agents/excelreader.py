#
# GAMS - General Algebraic Modeling System Python API
#
# Copyright (c) 2017-2024 GAMS Development Corp. <support@gams.com>
# Copyright (c) 2017-2024 GAMS Software GmbH <support@gams.com>
#
# Permission is hereby granted, free of charge, to any person obtaining a copy
# of this software and associated documentation files (the "Software"), to deal
# in the Software without restriction, including without limitation the rights
# to use, copy, modify, merge, publish, distribute, sublicense, and/or sell
# copies of the Software, and to permit persons to whom the Software is
# furnished to do so, subject to the following conditions:
#
# The above copyright notice and this permission notice shall be included in all
# copies or substantial portions of the Software.
#
# THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
# IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
# FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
# AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
# LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM,
# OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE
# SOFTWARE.
#

from copy import copy
import os
import sys
import datetime
import gams.transfer as gt
import openpyxl
from openpyxl.utils.cell import column_index_from_string
from pandas.api.types import is_datetime64_any_dtype as is_datetime
import pandas as pd
import numpy as np
from gams.core.gdx import GMS_SV_UNDEF
from gams.transfer.syms._methods.tables import (
    _assert_axes_no_nans,
    _get_implied_dimension_from_axes,
    _flatten_and_convert,
)
from gams.connect.agents.excelagent import ExcelAgent
from cerberus import Validator


class ExcelReader(ExcelAgent):
    _index_parameter_map = {
        "rdim": "rowDimension",
        "rowdimension": "rowDimension",
        "cdim": "columnDimension",
        "columndimension": "columnDimension",
        "skipempty": "skipEmpty",
        "se": "skipEmpty",
        "ignoretext": "ignoreText",
        "automerge": "autoMerge",
        "ignorerows": "ignoreRows",
        "ignorecolumns": "ignoreColumns",
        "mergedcells": "mergedCells",
    }

    def __init__(self, system_directory, cdb, inst):
        super().__init__(system_directory, cdb, inst)

        self._parse_options()

        if self._trace > 3:
            pd.set_option("display.max_rows", None, "display.max_columns", None)
            np.set_printoptions(threshold=sys.maxsize)
        if os.path.splitext(self._file)[1] in [".xls"]:
            self.connect_error("The ExcelReader does not support .xls files.")

    def _parse_options(self):
        self._file = self._inst["file"]
        self._cdim = self._inst.get("columnDimension", 1)
        self._rdim = self._inst.get("rowDimension", 1)
        self._sym_type = self._inst.get("type", "par")
        self._symbols = self._inst.get("symbols", [])
        self._merged_cells = self._inst.get("mergedCells", False)
        self._skip_empty = self._inst.get("skipEmpty", 1)
        self._value_subs = self._inst.get("valueSubstitutions", None)
        self._index_subs = self._inst.get("indexSubstitutions", None)
        self._index = self._inst.get("index", None)
        self._auto_merge = self._inst.get("autoMerge", False)
        self._ignore_text = self._inst.get("ignoreText", "auto")
        self._trace = self._inst.get("trace", 0)

        if self._trace > 1:
            self._cdb.print_log(
                "Input (root):"
                f"\n  file: >{self._inst['file']}<"
                f"\n  rowDimension: >{self._inst.get('rowDimension', '')}<"
                f"\n  columnDimension: >{self._inst.get('columnDimension', '')}<"
                f"\n  type: >{self._inst.get('type', '')}<"
                f"\n  mergedCells: >{self._inst.get('mergedCells', '')}<"
                f"\n  skipEmpty: >{self._inst.get('skipEmpty', '')}<"
                f"\n  valueSubstitutions: >{self._inst.get('valueSubstitutions', '')}<"
                f"\n  indexSubstitutions: >{self._inst.get('indexSubstitutions', '')}<"
                f"\n  index: >{self._inst.get('index', '')}<"
                f"\n  autoMerge: >{self._inst.get('autoMerge', '')}<"
                f"\n  ignoreText: >{self._inst.get('ignoreText', '')}<"
                f"\n  trace: >{self._inst.get('trace', '')}<"
                "\n"
            )

        self._file = os.path.abspath(self._file)

        if self._trace > 1:
            self._cdb.print_log(
                "Processed Input (root):"
                f"\n  file: >{self._file}<"
                f"\n  rowDimension: >{self._rdim}<"
                f"\n  columnDimension: >{self._cdim}<"
                f"\n  type: >{self._sym_type}<"
                f"\n  mergedCells: >{self._merged_cells}<"
                f"\n  skipEmpty: >{self._skip_empty}<"
                f"\n  valueSubstitutions: >{self._value_subs}<"
                f"\n  indexSubstitutions: >{self._index_subs}<"
                f"\n  index: >{self._index}<"
                f"\n  autoMerge: >{self._auto_merge}<"
                f"\n  ignoreText: >{self._ignore_text}<"
                f"\n  trace: >{self._trace}<"
                "\n"
            )

    def _apply_skip_empty(self, dim, idx, skip_empty):
        stop = None
        count = 0
        if dim > 0 and skip_empty > -1:
            for i in range(idx.shape[1]):
                if (np.array(idx[:, i] == None)).all():
                    count += 1
                else:
                    count = 0
                if count > skip_empty:
                    stop = i - skip_empty
                    break
        return stop

    def _create_index(self, dim, idx):
        if dim > 1:
            return pd.MultiIndex.from_arrays(idx)
        else:
            return idx.flatten()

    def _remove_missing_index(self, values, rdim, cdim, row_idx, col_idx, method):
        def _keep_list(idx):
            keep = list(range(idx.shape[1]))
            for i in reversed(range(idx.shape[1])):
                if method(v is None for v in idx[:, i]):
                    del keep[i]
            return keep

        if rdim > 0:
            keep = _keep_list(row_idx)
            row_idx = row_idx[:, keep]
            values = values[keep]

        if cdim > 0:
            keep = _keep_list(col_idx)
            col_idx = col_idx[:, keep]
            values = values[:, keep]

        return values, row_idx, col_idx

    def _apply_auto_merge(self, idx, dim):
        last_label = [None] * dim
        for i in range(idx.shape[1]):
            if any(idx[:, i] != None):
                for j in range(idx.shape[0]):
                    if idx[j, i] is None:
                        idx[j, i] = last_label[j]
                last_label = idx[:, i]
        return idx

    def _create_dataframe(self, col_idx, row_idx, values, rdim, cdim):
        # create column and row index used for DataFrame
        col_idx = self._create_index(cdim, col_idx)
        row_idx = self._create_index(rdim, row_idx)

        if cdim == rdim == 0:
            df = pd.DataFrame(values.flatten())
        elif cdim == 0:
            values = values[:, 0]
            df = pd.DataFrame(values.flatten(), index=row_idx)
        elif rdim == 0:
            values = values[0, :]
            df = pd.DataFrame(values.flatten(), index=col_idx)
        else:
            df = pd.DataFrame(values, index=row_idx, columns=col_idx)
        return df

    def _resolve_merged_cells(self, sheet, data):
        # TODO: do this only on the used range for better performance
        for mr in sheet.merged_cells.ranges:
            nwc, nwr, sec, ser = mr.bounds
            value = data[nwr - 1][nwc - 1]
            data[nwr - 1 : ser, nwc - 1 : sec] = value
        return data

    def _convert_dates(self, df):
        for col in df.columns:
            if is_datetime(df[col]):
                df[col] = (
                    pd.DatetimeIndex(df[col]).to_julian_date()
                    - pd.Timestamp("1899-12-30").to_julian_date()
                )

        has_datetime = any(
            isinstance(x, datetime.datetime) for x in df.values.flatten()
        )
        if has_datetime:
            if hasattr(pd.DataFrame, "map"):
                df = df.map(
                    lambda x: (
                        pd.Timestamp(x).to_julian_date()
                        - pd.Timestamp("1899-12-30").to_julian_date()
                        if isinstance(x, datetime.datetime)
                        else x
                    )
                )
            else:
                df = df.applymap(
                    lambda x: (
                        pd.Timestamp(x).to_julian_date()
                        - pd.Timestamp("1899-12-30").to_julian_date()
                        if isinstance(x, datetime.datetime)
                        else x
                    )
                )

        return df

    def _index_substitutions(self, row_idx, col_idx, rdim, cdim, index_subs):
        if index_subs and rdim + cdim > 0:
            subs = copy(index_subs)
            for k, v in subs.items():
                if k is None:
                    break
                if k != k:  # check for float('nan')
                    if None not in subs.keys():
                        subs[None] = v
                    break
            if rdim > 0:
                ri_tmp = copy(row_idx)
                for k, v in subs.items():
                    # change value in array if either row_idx==k or (element is .nan (row_idx!=row_idx) and key is .nan (k!=k))
                    ri_tmp[
                        np.logical_or(
                            row_idx == k, np.logical_and(row_idx != row_idx, k != k)
                        )
                    ] = v
                row_idx = ri_tmp
            if cdim > 0:
                ci_tmp = copy(col_idx)
                for k, v in subs.items():
                    # change value in array if either col_idx==k or (element is .nan (col_idx!=col_idx) and key is .nan (k!=k))
                    ci_tmp[
                        np.logical_or(
                            col_idx == k, np.logical_and(col_idx != col_idx, k != k)
                        )
                    ] = v
                col_idx = ci_tmp
        return row_idx, col_idx
        # alternative approach, but much slower
        # vectorized_replace_rows = np.vectorize(lambda v: dict().get(v, v), otypes=[row_idx.dtype])
        # vectorized_replace_cols = np.vectorize(lambda v: dict().get(v, v), otypes=[col_idx.dtype])
        #
        # if rdim > 0:
        #    row_idx = vectorized_replace_rows(row_idx)
        # if cdim > 0:
        #    col_idx = vectorized_replace_cols(col_idx)

    def _value_substitutions(self, df, value_sub):
        if value_sub is not None:
            # pandas-version-check
            if self.pandas_version_before(pd.__version__, "2.2"):  # pandas < 2.2.0
                df.replace(value_sub, inplace=True)
            else:  # pandas >= 2.2.0
                with pd.option_context("future.no_silent_downcasting", True):
                    df = df.replace(value_sub).infer_objects()
        return df

    def _write(self, df, sym_name, sym_type, rdim, cdim):
        if df is None or df.empty or df.isnull().all().all():
            df = None
        # pass DataFrame as Series for rdim=0 or cdim=0 to not confuse gams.transfer with dimensions
        elif (cdim == 0 and rdim != 0) or (rdim == 0 and cdim != 0):
            df = df[0]

        if sym_type == "par":
            sym = self._cdb._container.addParameter(
                sym_name,
                ["*"] * (rdim + cdim),
                records=df,
                uels_on_axes=True,
            )
            if df is not None:
                sym.dropUndef()  # drop float('nan')
                # TODO: remove this section as soon as gams.transfer supports dropping NaN values
                sym.records = self._value_substitutions(
                    sym.records, {GMS_SV_UNDEF: gt.SpecialValues.UNDEF}
                )

        else:  # set
            if df is not None:
                # TODO: remove this section as soon as gams.transfer supports dropping NaN values
                # Nan values become empty set element text and we can not drope those values after they are in the container.
                # This is the workaround to handle this
                _assert_axes_no_nans(df)
                dim = _get_implied_dimension_from_axes(df)
                if dim != rdim + cdim:
                    self.connect_error(
                        f"Dimensionality of table ({dim}) is inconsistent with set domain specification ({rdim+cdim})"
                    )
                df = _flatten_and_convert(df)
                df.dropna(inplace=True)
            sym = self._cdb._container.addSet(
                sym_name,
                ["*"] * (rdim + cdim),
                records=df,
                uels_on_axes=False,
            )

        # For symbols with None records, empty df is assigned
        self.transform_sym_none_to_empty(sym)

    def _ignore_rows(self, data, offset, ignore_rows):
        nr_rows = data.shape[0]
        r = list(
            filter(
                lambda x: x + offset + 1 not in ignore_rows,
                range(nr_rows),
            )
        )
        data = data[r]
        return data

    def _ignore_columns(self, data, offset, ignore_columns):
        nr_cols = data.shape[1]
        r = list(
            filter(
                lambda x: x + offset + 1 not in ignore_columns,
                range(nr_cols),
            )
        )
        data = data[:, r]
        return data

    def _apply_ignore_rows_columns(
        self, data, ignore_rows, ignore_columns, nw_row, nw_col, sym_name
    ):
        # apply ignoreRows
        if ignore_rows is not None:
            data = self._ignore_rows(data, nw_row, ignore_rows)
            if self._trace > 2:
                self._cdb.print_log(
                    f"Raw data after ignoreRows ({sym_name}):\n{data}\n"
                )

        # apply ignoreColumns
        if ignore_columns is not None:
            data = self._ignore_columns(data, nw_col, ignore_columns)
            if self._trace > 2:
                self._cdb.print_log(
                    f"Raw data after ignoreColumns ({sym_name}):\n{data}\n"
                )
        return data

    def _parse_ignore_rows(self, ignore_rows, nw_row, se_row):
        if ignore_rows is None:
            return []
        if isinstance(ignore_rows, int):
            l = [ignore_rows]
        else:
            l = ignore_rows
        l = set(l)
        if se_row is None:
            return list(l)
        l = list(filter(lambda x: x >= nw_row and x <= se_row, l))
        return l

    def _parse_ignore_columns(self, ignore_columns, nw_col, se_col):
        if ignore_columns is None:
            return []
        if isinstance(ignore_columns, int):
            l = [ignore_columns]
        elif isinstance(ignore_columns, str):
            l = [column_index_from_string(ignore_columns)]
        else:  # list
            l = []
            for c in ignore_columns:
                if isinstance(c, int):
                    l.append(c)
                else:  # string
                    l.append(column_index_from_string(c))
        l = set(l)
        if se_col is None:
            return list(l)
        l = list(filter(lambda x: x >= nw_col and x <= se_col, l))
        return l

    def _read_symbol(self, sym):
        if self._trace > 1:
            self._cdb.print_log(
                "Input (symbol):"
                f"\n  name: >{sym.get('name', '')}<"
                f"\n  range: >{sym.get('range', '')}<"
                f"\n  rowDimension: >{sym.get('rowDimension', '')}<"
                f"\n  columnDimension: >{sym.get('columnDimension', '')}<"
                f"\n  type: >{sym.get('type', '')}<"
                f"\n  mergedCells: >{sym.get('mergedCells', '')}<"
                f"\n  valueSubstitutions: >{sym.get('valueSubstitutions', '')}<"
                f"\n  indexSubstitutions: >{sym.get('indexSubstitutions', '')}<"
                f"\n  skipEmpty: >{sym.get('skipEmpty', '')}<"
                f"\n  autoMerge: >{sym.get('autoMerge', '')}<"
                f"\n  ignoreText: >{sym.get('ignoreText', '')}<"
                f"\n  ignoreRows: >{sym.get('ignoreRows', '')}<"
                f"\n  ignoreColumns: >{sym.get('ignoreColumns', '')}<"
                "\n"
            )
        sym_name = sym["name"]
        sym_range = sym.get("range", sym_name + "!A1")
        sym_type = sym.get("type", self._sym_type)
        merged_cells = sym.get("mergedCells", self._merged_cells)
        rdim = sym.get("rowDimension", self._rdim)
        cdim = sym.get("columnDimension", self._cdim)
        value_subs = sym.get("valueSubstitutions", self._value_subs)
        index_subs = sym.get("indexSubstitutions", self._index_subs)
        skip_empty = sym.get("skipEmpty", self._skip_empty)
        auto_merge = sym.get("autoMerge", self._auto_merge)
        ignore_text = sym.get("ignoreText", self._ignore_text)

        sheet, nw_col, nw_row, se_col, se_row, _ = self.parse_range(sym_range, self._wb)
        nw_only = se_col is None and se_row is None
        ignore_rows = self._parse_ignore_rows(sym.get("ignoreRows"), nw_row, se_row)
        ignore_columns = self._parse_ignore_columns(
            sym.get("ignoreColumns"), nw_col, se_col
        )
        required_rows = cdim + 1 + len(ignore_rows)
        required_cols = rdim + 1 + len(ignore_columns)
        if not nw_only:
            nr_cols = se_col - nw_col
            nr_rows = se_row - nw_row

        # handle ignoreText=auto
        if sym_type == "set" and ignore_text == "auto":
            ignore_text = False
            if rdim == 0:
                if (
                    nw_only or nr_rows < required_rows
                ):  # nw only or range without set element text
                    ignore_text = True
            if cdim == 0:
                if (
                    nw_only or nr_cols < required_cols
                ):  # nw only or range without set element text
                    ignore_text = True

        if self._trace > 1:
            self._cdb.print_log(
                "Processed Input (symbol):"
                f"\n  name: >{sym_name}<"
                f"\n  range: >{sym_range}<"
                f"\n  rowDimension: >{rdim}<"
                f"\n  columnDimension: >{cdim}<"
                f"\n  type: >{sym_type}<"
                f"\n  mergedCells: >{merged_cells}<"
                f"\n  valueSubstitutions: >{value_subs}<"
                f"\n  indexSubstitutions: >{index_subs}<"
                f"\n  skipEmpty: >{skip_empty}<"
                f"\n  autoMerge: >{auto_merge}<"
                f"\n  ignoreText: >{ignore_text}<"
                f"\n  ignoreRows: >{ignore_rows}<"
                f"\n  ignoreColumns: >{ignore_columns}<"
                "\n"
            )

        # check that sets do not have dim=0
        if sym_type == "set" and rdim == 0 and cdim == 0:
            self.connect_error(
                f"Cannot read set >{sym_name}< with both rowDimension=0 and columnDimension=0."
            )

        # check sufficient ranges
        if sym_type == "set" and ignore_text:
            if cdim == 0:
                required_cols -= 1
            elif rdim == 0:
                required_rows -= 1
        if not nw_only:
            if sym_type == "set" and not ignore_text:
                if cdim == 0 and nr_cols == required_cols - 1:
                    self.connect_error(
                        "Range and rowDimension specification does not contain set element text but ignoreText has been set to False. Adjust range or rowDimension or set ignoreText=True."
                    )
                if rdim == 0 and nr_rows == required_rows - 1:
                    self.connect_error(
                        "Range and columnDimension specification does not contain set element text but ignoreText has been set to False. Adjust range or columnDimension or set ignoreText=True."
                    )
            if nr_rows < required_rows:
                self.connect_error(
                    f"Invalid range >{sym_range}<. With columnDimension: >{cdim}< and {len(ignore_rows)} rows to be ignored, the range must include at least {required_rows} rows."
                )
            if nr_cols < required_cols:
                self.connect_error(
                    f"Invalid range >{sym_range}<. With rowDimension: >{rdim}< and {len(ignore_columns)} columns to be ignored, the range must include at least {required_cols} columns."
                )

        data = np.array(
            list(sheet.values), dtype=object
        )  # dtype=object is required to not convert int values (e.g. 1) to float automatically (e.g. 1.0)

        if len(data) == 0:  # no data at all
            self._write(None, sym_name, sym_type, rdim, cdim)
            return

        if self._trace > 2:
            self._cdb.print_log(f"Raw data ({sym_name}) :\n{data}\n")

        if merged_cells:
            data = self._resolve_merged_cells(sheet, data)
            if self._trace > 2:
                self._cdb.print_log(
                    f"Raw data after resolving merged cells ({sym_name}):\n{data}\n"
                )

        # shrink data to actual range
        data = data[nw_row:se_row, nw_col:se_col]
        if self._trace > 2:
            self._cdb.print_log(
                f"Raw data after shrinking to range ({sym_name}):\n{data}\n"
            )

        # apply ignoreRows and ignoreColumns
        data = self._apply_ignore_rows_columns(
            data, ignore_rows, ignore_columns, nw_row, nw_col, sym_name
        )
        if len(data) == 0:
            self._write(None, sym_name, sym_type, rdim, cdim)
            return
        # if data.shape[0] < required_rows - len(ignore_rows):
        #    self.connect_error(
        #        f"Insufficient number of data rows ({sym_name}). Require at least {required_rows}, but got {data.shape[0]}."
        #    )
        # if data.shape[1] < required_cols - len(ignore_columns):
        #    self.connect_error(
        #        f"Insufficient number of data columns ({sym_name}). Require at least {required_cols}, but got {data.shape[1]}."
        #    )

        col_idx = data[:cdim, rdim:]
        row_idx = data[cdim:, :rdim].transpose()

        if self._trace > 2:
            self._cdb.print_log(f"Initial column index ({sym_name}):\n{col_idx}\n")
            self._cdb.print_log(f"Initial row index ({sym_name}):\n{row_idx}\n")

        # apply skipEmpty only for nw_only, but not for explicit ranges
        if nw_only:
            stop_col = self._apply_skip_empty(cdim, col_idx, skip_empty)
            col_idx = col_idx[:, :stop_col]
            stop_row = self._apply_skip_empty(rdim, row_idx, skip_empty)
            row_idx = row_idx[:, :stop_row]
            if self._trace > 2:
                self._cdb.print_log(
                    f"Column index after skipEmpty ({sym_name}):\n{col_idx}\n"
                )
                self._cdb.print_log(
                    f"Row index after skipEmpty ({sym_name}):\n{row_idx}\n"
                )
        else:
            stop_col = None
            stop_row = None

        if stop_col is not None:
            stop_col += rdim
        if stop_row is not None:
            stop_row += cdim

        if cdim == 0 and rdim == 0:  # handle scalars
            stop_row = 1
            stop_col = 1
        if rdim == 0:
            row_idx = np.empty((0, 0))  # dummy array for header
        if cdim == 0:
            col_idx = np.empty((0, 0))  # dummy array for header

        values = data[cdim:stop_row, rdim:stop_col]

        if self._trace > 2:
            self._cdb.print_log(f"Values {(sym_name)}: {values}\n")

        if auto_merge:
            if cdim > 1:
                col_idx = self._apply_auto_merge(col_idx, cdim)
            if rdim > 1:
                row_idx = self._apply_auto_merge(row_idx, rdim)
            if self._trace > 2:
                self._cdb.print_log(
                    f"Row index after autoMerge ({sym_name}):\n{row_idx}\n"
                )
                self._cdb.print_log(
                    f"Column index after autoMerge ({sym_name}):\n{col_idx}\n"
                )

        # replace all set text with empty string for ignoreText=True
        if sym_type == "set" and ignore_text:
            if values.size == 0:
                if cdim == 0:
                    values = np.empty((values.shape[0], 1), dtype=str)
                elif rdim == 0:
                    values = np.empty((1, values.shape[1]), dtype=str)
            else:
                values = np.empty_like(values, dtype=str)

        if index_subs:
            # remove all-None entries in column and row header and corresponding values
            values, row_idx, col_idx = self._remove_missing_index(
                values, rdim, cdim, row_idx, col_idx, all
            )

            row_idx, col_idx = self._index_substitutions(
                row_idx, col_idx, rdim, cdim, index_subs
            )
            if self._trace > 2:
                self._cdb.print_log(
                    f"Row index after indexSubstitutions ({sym_name}):\n{row_idx}\n"
                )
                self._cdb.print_log(
                    f"Column index after indexSubstitutions ({sym_name}):\n{col_idx}\n"
                )

        # remove any-None entries in column and row header and corresponding values
        values, row_idx, col_idx = self._remove_missing_index(
            values, rdim, cdim, row_idx, col_idx, any
        )

        if self._trace > 2:
            self._cdb.print_log(
                f"Column index before DataFrame creation ({sym_name}):\n{col_idx}\n"
            )
            self._cdb.print_log(
                f"Row index before DataFrame creation ({sym_name}):\n{row_idx}\n"
            )
            self._cdb.print_log(
                f"Values before DataFrame creation ({sym_name}):\n{values}\n"
            )

        df = self._create_dataframe(col_idx, row_idx, values, rdim, cdim)

        if self._trace > 2:
            self._cdb.print_log(f"Initial DataFrame ({sym_name}):\n{df}\n")

        df = self._convert_dates(df)

        df = self._value_substitutions(df, value_subs)
        if self._trace > 2:
            self._cdb.print_log(
                f"DataFrame after valueSubstitutions ({sym_name}):\n{df}\n"
            )

        # TODO: This is a workaround to get UNDEF to survive sym.dropNA/sym.dropUndef - remove as soon as gams.transfer supports dropping NaN values
        if sym_type == "par":
            import re

            pattern = re.compile(r"undef", re.IGNORECASE)
            # pandas-version-check
            if self.pandas_version_before(pd.__version__, "2.2"):  # pandas < 2.2.0
                df.replace(regex=pattern, value=GMS_SV_UNDEF, inplace=True)
            else:  # pandas >= 2.2.0
                with pd.option_context("future.no_silent_downcasting", True):
                    df = df.replace(regex=pattern, value=GMS_SV_UNDEF).infer_objects()
        self._write(df, sym_name, sym_type, rdim, cdim)

    def open(self):
        read_only = not (
            any(sym.get("mergedCells", self._merged_cells) for sym in self._symbols)
            or self._merged_cells
        )
        self._wb = openpyxl.load_workbook(
            self._file, read_only=read_only, data_only=True
        )  # data_only=True is required to read values instead of formulas

    def _read_symbols(self, symbols, validate=False):
        for i, sym in enumerate(symbols):
            if validate:
                sym_schema = self.cerberus()["symbols"]["schema"]["schema"]
                v = Validator(sym_schema)
                if not v.validate(sym):
                    self.connect_error(
                        f"Validation of item {i} in index failed: {v.errors}"
                    )
            self._read_symbol(sym)

    def _create_symbol_instructions(self, rec):
        is_symbol = not None in (rec[0], rec[1], rec[2])
        inst = {}
        if is_symbol:
            inst["type"] = rec[0].lower().strip()
            inst["name"] = rec[1].strip()
            inst["range"] = rec[2].strip()
            if self._trace > 1:
                self._cdb.print_log(
                    f"\nIndex sheet: Parse symbol >{inst['name']}< with type=>{inst['type']}< and range=>{inst['range']}<."
                )
        return inst

    def _finalize_symbol_instructions(self, instructions):
        instructions["rowDimension"] = instructions.get("rowDimension", 0)
        instructions["columnDimension"] = instructions.get("columnDimension", 0)
        return instructions

    def _read_from_index(self):
        symbols = self.parse_index(self._index, self._wb, self._index_parameter_map)

        # reopen the file with read_only=False if required
        if not self._merged_cells:
            read_only = not any(
                sym.get("mergedCells", self._merged_cells) for sym in symbols
            )
            if not read_only:
                self._wb.close()
                self._wb = openpyxl.load_workbook(
                    self._file, read_only=read_only, data_only=True
                )  # data_only=True is required to read values instead of formulas
        self._read_symbols(symbols, True)

    def execute(self):
        if self._index:
            self._read_from_index()
        else:
            self._read_symbols(self._symbols, True)
        if self._trace > 0:
            self.describe_container(self._cdb._container, "Connect Container")
        if self._trace > 2:
            for name, sym in self._cdb._container.data.items():
                self._cdb.print_log(
                    f"Connect Container symbol >{name}<:\n {sym.records}\n"
                )

    def close(self):
        self._wb.close()
